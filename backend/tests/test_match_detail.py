"""Tests for AccountMatchDetail preservation and exact linking"""
import pytest
from datetime import datetime
from openpyxl import load_workbook

from app.domain.models import (
    InstagramPost, PostType, AuditResult, PostAuditStatus,
    AccountMatchDetail, MatchStatus
)
from app.exporters.excel_exporter import ExcelExporter


@pytest.fixture
def multi_match_audit():
    """MASTER_A/B with multiple reviews - exact link validation"""
    master_a = InstagramPost(
        shortcode="MASTER_A", username="master", caption="Post A",
        post_type=PostType.PHOTO, published_at=datetime(2026, 9, 1),
        permalink="https://www.instagram.com/p/MASTER_A/", is_video=False,
    )
    master_b = InstagramPost(
        shortcode="MASTER_B", username="master", caption="Post B",
        post_type=PostType.PHOTO, published_at=datetime(2026, 9, 2),
        permalink="https://www.instagram.com/p/MASTER_B/", is_video=False,
    )

    return AuditResult(
        master_account="master",
        compared_accounts=["account_a", "account_b"],
        posts_status=[
            PostAuditStatus(
                master_post=master_a,
                account_status={"account_a": "review", "account_b": "review"},
                account_matches={
                    "account_a": AccountMatchDetail(
                        status=MatchStatus.REVIEW, score=0.75,
                        candidate_shortcode="A1",
                        candidate_permalink="https://www.instagram.com/p/A1/",
                        visual_similarity=0.7, caption_similarity=0.75,
                        type_match=True, date_similarity=0.8,
                    ),
                    "account_b": AccountMatchDetail(
                        status=MatchStatus.REVIEW, score=0.72,
                        candidate_shortcode="A2",
                        candidate_permalink="https://www.instagram.com/p/A2/",
                        visual_similarity=0.65, caption_similarity=0.72,
                        type_match=True, date_similarity=0.75,
                    ),
                },
                missing_in=[], review_in=["account_a", "account_b"],
            ),
            PostAuditStatus(
                master_post=master_b,
                account_status={"account_a": "review", "account_b": "missing"},
                account_matches={
                    "account_a": AccountMatchDetail(
                        status=MatchStatus.REVIEW, score=0.68,
                        candidate_shortcode="B1",
                        candidate_permalink="https://www.instagram.com/p/B1/",
                        visual_similarity=0.6, caption_similarity=0.68,
                        type_match=True, date_similarity=0.7,
                    ),
                    "account_b": AccountMatchDetail(
                        status=MatchStatus.MISSING, score=0.3,
                        candidate_shortcode=None, candidate_permalink=None,
                        visual_similarity=0.2, caption_similarity=0.3,
                        type_match=False, date_similarity=0.1,
                    ),
                },
                missing_in=["account_b"], review_in=["account_a"],
            ),
        ],
        total_master_posts=2, found_in_all=0, with_missing=1, with_review=1,
        missing_count_by_account={"account_a": 0, "account_b": 1},
        review_count_by_account={"account_a": 2, "account_b": 1},
        timestamp=datetime.now(),
    )


def test_review_preserves_candidate_permalink(multi_match_audit):
    """REVIEW preserves candidate_permalink"""
    match = multi_match_audit.posts_status[0].account_matches["account_a"]
    assert match.status == MatchStatus.REVIEW
    assert match.candidate_permalink == "https://www.instagram.com/p/A1/"


def test_missing_hides_candidate_permalink(multi_match_audit):
    """MISSING has no candidate_permalink"""
    match = multi_match_audit.posts_status[1].account_matches["account_b"]
    assert match.status == MatchStatus.MISSING
    assert match.candidate_permalink is None


def test_revisar_multiple_reviews_same_master_keep_correct_links(multi_match_audit, tmp_path):
    """3_REVISAR: MASTER_A/account_a, account_b have correct links (exact)"""
    exporter = ExcelExporter()
    result_path = exporter.export(multi_match_audit, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["3_REVISAR"]

    # Row 2: MASTER_A/account_a
    assert ws.cell(2, 2).hyperlink.target == "https://www.instagram.com/p/MASTER_A/"
    assert ws.cell(2, 7).hyperlink.target == "https://www.instagram.com/p/A1/"

    # Row 3: MASTER_A/account_b
    assert ws.cell(3, 2).hyperlink.target == "https://www.instagram.com/p/MASTER_A/"
    assert ws.cell(3, 7).hyperlink.target == "https://www.instagram.com/p/A2/"


def test_revisar_multiple_master_posts_do_not_cross_links(multi_match_audit, tmp_path):
    """3_REVISAR: MASTER_B/account_a has correct links (exact), no cross-link"""
    exporter = ExcelExporter()
    result_path = exporter.export(multi_match_audit, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["3_REVISAR"]

    # Row 4: MASTER_B/account_a
    assert ws.cell(4, 2).hyperlink.target == "https://www.instagram.com/p/MASTER_B/"
    assert ws.cell(4, 7).hyperlink.target == "https://www.instagram.com/p/B1/"


def test_parrilla_found_links_candidate_exactly(multi_match_audit, tmp_path):
    """2_PARRILLA FOUND links to exact candidate_permalink"""
    fixture_with_found = AuditResult(
        master_account="master",
        compared_accounts=["account_a"],
        posts_status=[
            PostAuditStatus(
                master_post=InstagramPost(
                    shortcode="FOUND_POST", username="master", caption="Found post",
                    post_type=PostType.PHOTO, published_at=datetime(2026, 9, 1),
                    permalink="https://www.instagram.com/p/FOUND_POST/", is_video=False,
                ),
                account_status={"account_a": "found"},
                account_matches={
                    "account_a": AccountMatchDetail(
                        status=MatchStatus.FOUND, score=1.0,
                        candidate_shortcode="FOUND_A", candidate_permalink="https://www.instagram.com/p/FOUND_A/",
                        visual_similarity=1.0, caption_similarity=1.0, type_match=True, date_similarity=1.0,
                    ),
                },
                missing_in=[], review_in=[],
            ),
        ],
        total_master_posts=1, found_in_all=1, with_missing=0, with_review=0,
        missing_count_by_account={"account_a": 0}, review_count_by_account={"account_a": 0},
        timestamp=datetime.now(),
    )
    exporter = ExcelExporter()
    result_path = exporter.export(fixture_with_found, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["2_PARRILLA"]

    # Row 2: FOUND_POST/account_a should link to candidate
    assert ws.cell(2, 2).hyperlink.target == "https://www.instagram.com/p/FOUND_A/"


def test_parrilla_review_links_candidate_exactly(multi_match_audit, tmp_path):
    """2_PARRILLA REVIEW links to exact candidate_permalink"""
    exporter = ExcelExporter()
    result_path = exporter.export(multi_match_audit, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["2_PARRILLA"]

    # Row 2: MASTER_A with account_a (col 2) and account_b (col 3)
    assert ws.cell(2, 2).hyperlink.target == "https://www.instagram.com/p/A1/"
    # Row 2: MASTER_A with account_b (col 3)
    assert ws.cell(2, 3).hyperlink.target == "https://www.instagram.com/p/A2/"


def test_parrilla_missing_links_original_exactly(multi_match_audit, tmp_path):
    """2_PARRILLA MISSING links to master original"""
    exporter = ExcelExporter()
    result_path = exporter.export(multi_match_audit, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["2_PARRILLA"]

    # Row 3: MASTER_B with account_a (col 2) REVIEW and account_b (col 3) MISSING
    assert ws.cell(3, 2).hyperlink.target == "https://www.instagram.com/p/B1/"
    # Row 3: MASTER_B account_b MISSING → master original
    assert ws.cell(3, 3).hyperlink.target == "https://www.instagram.com/p/MASTER_B/"


def test_escaneo_found_links_candidate_exactly(multi_match_audit, tmp_path):
    """1_ESCANEO FOUND links to exact candidate_permalink"""
    fixture_with_found = AuditResult(
        master_account="master",
        compared_accounts=["account_a"],
        posts_status=[
            PostAuditStatus(
                master_post=InstagramPost(
                    shortcode="FOUND_POST", username="master", caption="Found post",
                    post_type=PostType.PHOTO, published_at=datetime(2026, 9, 1),
                    permalink="https://www.instagram.com/p/FOUND_POST/", is_video=False,
                ),
                account_status={"account_a": "found"},
                account_matches={
                    "account_a": AccountMatchDetail(
                        status=MatchStatus.FOUND, score=1.0,
                        candidate_shortcode="FOUND_A", candidate_permalink="https://www.instagram.com/p/FOUND_A/",
                        visual_similarity=1.0, caption_similarity=1.0, type_match=True, date_similarity=1.0,
                    ),
                },
                missing_in=[], review_in=[],
            ),
        ],
        total_master_posts=1, found_in_all=1, with_missing=0, with_review=0,
        missing_count_by_account={"account_a": 0}, review_count_by_account={"account_a": 0},
        timestamp=datetime.now(),
    )
    exporter = ExcelExporter()
    result_path = exporter.export(fixture_with_found, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["1_ESCANEO"]

    # Row 2: FOUND_POST master link
    assert ws.cell(2, 2).hyperlink.target == "https://www.instagram.com/p/FOUND_POST/"
    # Column 5 (account_a): FOUND status links to candidate
    assert ws.cell(2, 5).hyperlink.target == "https://www.instagram.com/p/FOUND_A/"


def test_escaneo_review_links_candidate_exactly(multi_match_audit, tmp_path):
    """1_ESCANEO REVIEW links to exact candidate_permalink"""
    exporter = ExcelExporter()
    result_path = exporter.export(multi_match_audit, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["1_ESCANEO"]

    # Row 2: MASTER_A/account_a (REVIEW) candidate link
    assert ws.cell(2, 5).hyperlink.target == "https://www.instagram.com/p/A1/"
    # Row 2: MASTER_A/account_b (REVIEW) candidate link
    assert ws.cell(2, 6).hyperlink.target == "https://www.instagram.com/p/A2/"


def test_escaneo_missing_has_no_link(multi_match_audit, tmp_path):
    """1_ESCANEO MISSING has no hyperlink"""
    exporter = ExcelExporter()
    result_path = exporter.export(multi_match_audit, str(tmp_path))
    wb = load_workbook(result_path)
    ws = wb["1_ESCANEO"]

    # Row 3: MASTER_B/account_b (MISSING) should have no hyperlink
    assert ws.cell(3, 6).hyperlink is None
