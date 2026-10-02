"""Excel exporter for Instagram content audit results"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from app.domain.models import AuditResult


class ExcelExporter:
    """Export audit results to Excel format"""

    def __init__(self, filename: str = "NewBody_Auditoria_Contenido.xlsx"):
        self.filename = filename

    def export(self, audit_result: AuditResult, output_dir: str = "data/exports") -> Path:
        """Export audit result to Excel file"""
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        file_path = output_path / self.filename
        workbook = Workbook()
        workbook.remove(workbook.active)

        # Create sheets
        self._create_escaneo_sheet(workbook, audit_result)
        self._create_faltantes_sheet(workbook, audit_result)
        self._create_revisar_sheet(workbook, audit_result)

        workbook.save(file_path)
        return file_path

    def _create_escaneo_sheet(self, workbook: Workbook, audit_result: AuditResult) -> None:
        """Create ESCANEO sheet with all posts"""
        ws = workbook.create_sheet("1_ESCANEO")

        # Headers
        headers = ["FECHA ORIGINAL", "LINK COLOMBIA", "CAPTION", "TIPO"]
        for account in audit_result.compared_accounts:
            headers.append(account.upper())
        headers.extend(["FALTA EN", "REVISAR EN"])

        ws.append(headers)

        # Format headers
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add data rows
        for post_status in audit_result.posts_status:
            master = post_status.master_post
            row = [
                master.published_at.strftime("%Y-%m-%d"),
                master.permalink,
                master.caption,
                master.post_type.value,
            ]

            # Add status for each account (ESTÁ, REVISAR, FALTA)
            for account in audit_result.compared_accounts:
                status = post_status.account_status.get(account, "missing")
                if status == "found":
                    row.append("ESTÁ")
                elif status == "review":
                    row.append("REVISAR")
                else:
                    row.append("FALTA")

            row.append(", ".join(post_status.missing_in) if post_status.missing_in else "")
            row.append(", ".join(post_status.review_in) if post_status.review_in else "")

            ws.append(row)

        # Format columns
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 12

        for account in audit_result.compared_accounts:
            ws.column_dimensions[get_column_letter(len(headers) - 2)].width = 15

        # Wrap text for caption column
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Convert URLs to hyperlinks
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")

        # Freeze panes
        ws.freeze_panes = "A2"

        # Auto filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    def _create_faltantes_sheet(self, workbook: Workbook, audit_result: AuditResult) -> None:
        """Create FALTANTES sheet with only missing posts"""
        ws = workbook.create_sheet("2_FALTANTES")

        # Headers
        headers = ["FECHA", "LINK COLOMBIA", "CAPTION", "TIPO", "PUBLICAR EN"]
        ws.append(headers)

        # Format headers
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add only posts with missing accounts
        for post_status in audit_result.posts_status:
            if not post_status.missing_in:
                continue

            master = post_status.master_post
            row = [
                master.published_at.strftime("%Y-%m-%d"),
                master.permalink,
                master.caption,
                master.post_type.value,
                ", ".join(post_status.missing_in),
            ]

            ws.append(row)

        # Format columns
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 30

        # Wrap text for caption column
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Convert URLs to hyperlinks
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")

        # Freeze panes
        ws.freeze_panes = "A2"

        # Auto filter
        ws.auto_filter.ref = f"A1:E1"

    def _create_revisar_sheet(self, workbook: Workbook, audit_result: AuditResult) -> None:
        """Create REVISAR sheet with only posts requiring review"""
        ws = workbook.create_sheet("3_REVISAR")

        # Headers
        headers = ["FECHA", "LINK COLOMBIA", "CAPTION", "TIPO", "REVISAR EN"]
        ws.append(headers)

        # Format headers
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add only posts with review accounts
        for post_status in audit_result.posts_status:
            if not post_status.review_in:
                continue

            master = post_status.master_post
            row = [
                master.published_at.strftime("%Y-%m-%d"),
                master.permalink,
                master.caption,
                master.post_type.value,
                ", ".join(post_status.review_in),
            ]

            ws.append(row)

        # Format columns
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 30

        # Wrap text for caption column
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Convert URLs to hyperlinks
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")

        # Freeze panes
        ws.freeze_panes = "A2"

        # Auto filter
        ws.auto_filter.ref = f"A1:E1"
